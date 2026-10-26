# -*- coding: utf-8 -*-
"""전체 흐름 시험 — LLM 대신 정해진 답을 돌려주는 가짜 provider 를 쓴다(네트워크 없음)."""
import io
import os
import re

import pytest

from unslop import pipeline
from unslop.formats import markdown_io
from unslop.units import parse_lines

FIX = {
    "선행연구 190편을 받아서 열고 본문을 읽었다": "선행연구 190편 조사 및 연구완료",
    "검증 손실은 좋아졌고 판정 능력은 나빠졌다": "검증 손실 개선, 판정 능력 하락",
}


def fake_provider(answers):
    def call(system, user):
        assert "출력 형식" in system
        out = []
        for uid, (kind, body) in parse_lines(user.split("# 줄별 지적 사항")[0]).items():
            plain = body.replace("⟦", "").replace("⟧", "")
            out.append("[%d|%s] %s" % (uid, kind, answers.get(plain, plain)))
        return "\n".join(out)
    return call


def test_markdown_roundtrip(tmp_path):
    src = tmp_path / "doc.md"
    src.write_text("# 제목입니다\n\n- 항목 하나\n\n| 구분 | 내용 |\n|---|---|\n| 데이터 | 12.95M 토큰 |\n\n```\n코드 줄\n```\n", encoding="utf-8")
    units, ctx = markdown_io.extract(str(src))
    assert [u.kind for u in units] == ["heading", "bullet", "cell", "cell", "cell", "cell"]
    out = tmp_path / "out.md"
    markdown_io.write(ctx, units, str(out))
    assert out.read_text(encoding="utf-8") == src.read_text(encoding="utf-8")


def test_run_brief_slide_markdown(tmp_path):
    src = tmp_path / "note.md"
    src.write_text("- 선행연구 190편을 받아서 열고 본문을 읽었다\n- 검증 손실은 좋아졌고 판정 능력은 나빠졌다\n- v4 는 JSON 출력 0/16\n",
                   encoding="utf-8")
    prep = pipeline.prepare(str(src), register="slide", workdir=str(tmp_path / "w"))
    assert prep["fixes"]                      # 「v4 는」 자동 수정
    acc = pipeline.rewrite(prep, fake_provider(FIX), rounds=2, log=lambda *a: None)
    rep = pipeline.finalize(prep, acc, str(tmp_path / "out.md"))
    text = (tmp_path / "out.md").read_text(encoding="utf-8")
    assert "선행연구 190편 조사 및 연구완료" in text and "v4는 JSON 출력 0/16" in text
    assert rep["after"]["errors"] == 0
    assert os.path.exists(os.path.join(rep["workdir"], "report.md"))


def test_gate_rejects_bad_rewrite(tmp_path):
    src = tmp_path / "s.md"
    src.write_text("- 학습데이터 12.95M 토큰 정리 완료\n", encoding="utf-8")
    prep = pipeline.prepare(str(src), register="slide", workdir=str(tmp_path / "w"))
    acc = pipeline.rewrite(prep, fake_provider({"학습데이터 12.95M 토큰 정리 완료": "학습데이터 13M 토큰 정리 완료"}),
                           rounds=2, log=lambda *a: None)
    assert acc == {}                          # 수치가 바뀌어 게이트에서 막힘 → 원문 유지
    assert prep["rejected"]


def test_apply_final_from_agent(tmp_path):
    src = tmp_path / "a.md"
    src.write_text("# 3. 학습 결과는 좋지 않았다\n\n데이터를 정제했다.\n", encoding="utf-8")
    prep = pipeline.prepare(str(src), register="report", workdir=str(tmp_path / "w"))
    final = tmp_path / "final.txt"
    final.write_text("[1|heading] 3. 학습 결과와 한계\n[2|para] 데이터를 정제했다.\n", encoding="utf-8")
    rep = pipeline.apply_final(prep["workdir"], str(final), str(tmp_path / "o.md"))
    assert (tmp_path / "o.md").read_text(encoding="utf-8").startswith("# 3. 학습 결과와 한계")
    assert rep["after"]["warnings"] == 0


def test_source_apply(tmp_path):
    from unslop.formats import source_io
    gen = tmp_path / "gen.py"
    gen.write_text('s = slide("검증 손실은 좋아졌고 판정 능력은 나빠졌다")\np = "D:\\\\usai\\\\papers 에 있다"\n', encoding="utf-8")
    r = source_io.apply(str(gen), [(1, "검증 손실은 좋아졌고 판정 능력은 나빠졌다", "검증 손실 개선, 판정 능력 하락"),
                                   (2, "D:\\usai\\papers 에 있다", "D:\\usai\\papers 위치")])
    t = gen.read_text(encoding="utf-8")
    assert "검증 손실 개선, 판정 능력 하락" in t and "D:\\\\usai\\\\papers 위치" in t
    assert r["changed"] == [1, 2]


def test_pptx_roundtrip(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches, Pt
    prs = pptx.Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(1))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = "선행연구 190편을 받아서 열고 본문을 읽었다"
    r.font.size = Pt(28)
    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    prep = pipeline.prepare(str(path), workdir=str(tmp_path / "w"))
    assert prep["register"] == "slide" and prep["units"][0].kind == "title"
    acc = pipeline.rewrite(prep, fake_provider(FIX), log=lambda *a: None)
    out = tmp_path / "deck2.pptx"
    pipeline.finalize(prep, acc, str(out))
    got = pptx.Presentation(str(out)).slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert got.text == "선행연구 190편 조사 및 연구완료" and got.font.size == Pt(28)


def test_docx_roundtrip(tmp_path):
    docx = pytest.importorskip("docx")
    d = docx.Document()
    d.add_heading("3. 학습 결과는 좋지 않았다", level=1)
    d.add_paragraph("데이터를 정제했다.")
    path = tmp_path / "r.docx"
    d.save(str(path))
    prep = pipeline.prepare(str(path), workdir=str(tmp_path / "w"))
    assert prep["register"] == "report" and prep["units"][0].kind == "heading"
    acc = pipeline.rewrite(prep, fake_provider({"3. 학습 결과는 좋지 않았다": "3. 학습 결과와 한계"}), log=lambda *a: None)
    out = tmp_path / "r2.docx"
    pipeline.finalize(prep, acc, str(out))
    assert docx.Document(str(out)).paragraphs[0].text == "3. 학습 결과와 한계"
