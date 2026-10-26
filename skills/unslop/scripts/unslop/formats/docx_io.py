"""DOCX 단위 추출과 반영 (python-docx 필요).

단위 = 문단 하나(본문·표 칸 안 문단 포함). 종류는 스타일 이름으로 정한다(Heading → heading, List → bullet).
반영할 때 글자 조각(run)이 하나면 그대로 바꾸고, 여러 개면 첫 조각에 전부 넣고 나머지를 비운다
— 이때 문단 안의 부분 서식(일부만 굵게 등)이 사라지므로 결과 보고에 개수를 남긴다.
"""
from __future__ import annotations

from typing import Dict, List, Tuple

from ..units import Unit, hangul_count


def _need():
    try:
        import docx  # noqa: F401
    except ImportError as e:  # pragma: no cover
        raise SystemExit("DOCX 처리에는 python-docx 가 필요합니다: pip install 'unslop[docx]'") from e


def _paragraphs(doc):
    for i, p in enumerate(doc.paragraphs):
        yield ("p", i), p, False
    for ti, t in enumerate(doc.tables):
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                for pi, p in enumerate(cell.paragraphs):
                    yield ("t", ti, ri, ci, pi), p, True


def _kind(p, in_table) -> str:
    if in_table:
        return "cell"
    name = (p.style.name if p.style is not None else "") or ""
    if name.lower().startswith(("heading", "title")) or name.startswith("제목"):
        return "heading"
    if "list" in name.lower() or "목록" in name:
        return "bullet"
    return "para"


def extract(path: str, min_hangul: int = 2) -> Tuple[List[Unit], Dict]:
    _need()
    import docx
    doc = docx.Document(path)
    units: List[Unit] = []
    for key, p, in_tab in _paragraphs(doc):
        t = p.text.strip()
        if hangul_count(t) >= min_hangul:
            units.append(Unit(len(units) + 1, _kind(p, in_tab), t, {"key": list(key)}))
    return units, {}


def write(ctx: Dict, units: List[Unit], out: str) -> Dict:
    _need()
    import docx
    doc = docx.Document(ctx["path"])
    want = {tuple(u.loc["key"]): u.text for u in units if ctx["original"].get(u.id) != u.text}
    changed = collapsed = 0
    for key, p, _ in _paragraphs(doc):
        new = want.get(tuple(key))
        if new is None:
            continue
        runs = p.runs
        if not runs:
            p.add_run(new)
        else:
            if len(runs) > 1:
                collapsed += 1
            runs[0].text = new
            for r in runs[1:]:
                r.text = ""
        changed += 1
    doc.save(out)
    return {"written": out, "paragraphs_changed": changed, "inline_format_collapsed": collapsed}
