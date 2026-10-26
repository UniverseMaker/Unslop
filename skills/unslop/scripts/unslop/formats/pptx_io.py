"""PPTX 단위 추출과 반영 (python-pptx 필요).

단위 = 글자 조각(run) 하나. 글자 크기·굵기·위치·표 여부로 종류를 정한다.
  title    20pt 이상
  subtitle 슬라이드 위쪽(1.6인치 안) 12~14pt 보통 굵기
  label    굵은 글자 (상자 제목·글머리 머리·태그)
  cell     표 칸 (머리 행 포함)
  body     그 밖
반영은 run.text 만 바꾸므로 글꼴·색·크기가 그대로 남는다. 같은 문구가 여러 곳에 있으면 모두 같은 단위로 묶는다.
"""
from __future__ import annotations

from typing import Dict, List, Tuple

from ..units import Unit, hangul_count


def _need():
    try:
        from pptx import Presentation  # noqa: F401
        from pptx.util import Inches  # noqa: F401
    except ImportError as e:  # pragma: no cover
        raise SystemExit("PPTX 처리에는 python-pptx 가 필요합니다: pip install 'unslop[pptx]'") from e


def _kind(run, top, in_table) -> str:
    from pptx.util import Inches
    if in_table:
        return "cell"
    size = run.font.size.pt if run.font.size else 12
    if size >= 20:
        return "title"
    if run.font.bold:
        return "label"
    if top is not None and top < Inches(1.6) and 12 <= size <= 14:
        return "subtitle"
    return "body"


def _runs(prs):
    for si, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            frames = []
            if sh.has_text_frame:
                frames.append((sh.text_frame, False))
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        frames.append((c.text_frame, True))
            for f, in_tab in frames:
                for p in f.paragraphs:
                    for r in p.runs:
                        yield si, sh, r, in_tab


def extract(path: str, min_hangul: int = 2) -> Tuple[List[Unit], Dict]:
    _need()
    from pptx import Presentation
    prs = Presentation(path)
    units: List[Unit] = []
    seen: Dict[str, Unit] = {}
    for si, sh, r, in_tab in _runs(prs):
        t = r.text.strip()
        if hangul_count(t) < min_hangul:
            continue
        if t in seen:
            seen[t].loc["slides"].append(si)
            continue
        u = Unit(len(units) + 1, _kind(r, sh.top, in_tab), t, {"slides": [si]})
        seen[t] = u
        units.append(u)
    return units, {}


def write(ctx: Dict, units: List[Unit], out: str) -> Dict:
    _need()
    from pptx import Presentation
    prs = Presentation(ctx["path"])
    orig = ctx["original"]          # {번호: 원문}
    m = {orig[u.id]: u.text for u in units if orig.get(u.id) is not None and orig[u.id] != u.text}
    hit = 0
    for _, _, r, _ in _runs(prs):
        t = r.text.strip()
        if t in m:
            r.text = r.text.replace(t, m[t])
            hit += 1
    prs.save(out)
    return {"written": out, "runs_changed": hit}
